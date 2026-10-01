"""
analysis/validate_spike.py

Reproducible HW4 skill extraction validation script.
Compares manual annotations in analysis/spike_validation.xlsx against
production extract_skills_from_posting results on analysis/spike_sample.json.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl

# Ensure src module is discoverable
ROOT_DIR = Path(__file__).resolve().parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from src.analytics import SkillVocabulary, extract_skills_from_posting

SAMPLE_PATH = ROOT_DIR / "analysis" / "spike_sample.json"
WORKBOOK_PATH = ROOT_DIR / "analysis" / "spike_validation.xlsx"
VOCAB_PATH = ROOT_DIR / "data" / "fixtures" / "skills_vocabulary.json"


def run_validation(
    sample_path: Path = SAMPLE_PATH,
    workbook_path: Path = WORKBOOK_PATH,
    vocab_path: Path = VOCAB_PATH,
) -> Dict[str, Any]:
    """
    Executes the manual annotation vs. production extractor validation.
    Returns calculated overall metrics, per-skill metrics, and itemized discrepancies.
    """
    with open(sample_path, "r", encoding="utf-8") as f:
        sample = json.load(f)

    vocab = SkillVocabulary.load_from_file(vocab_path)
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    ws = wb["Annotation"]

    # Read canonical skill names from header row (columns 9 to 60)
    skills = [ws.cell(1, c).value for c in range(9, 61)]

    total_postings = len(sample)
    total_decisions = 0
    evaluable_decisions = 0
    u_decisions = 0
    u_postings = set()

    per_skill: Dict[str, Dict[str, int]] = {
        s: {"human_pos": 0, "ext_pos": 0, "tp": 0, "fp": 0, "fn": 0, "tn": 0}
        for s in skills
    }

    discrepancies: List[Dict[str, Any]] = []

    for i, posting in enumerate(sample):
        sid = i + 1
        r = i + 2
        title = posting.get("title", "")
        company = posting.get("company", "")
        extracted = extract_skills_from_posting(posting, vocab)

        for c_idx, skill in enumerate(skills):
            c = c_idx + 9
            val = ws.cell(r, c).value
            total_decisions += 1

            # Handle unevaluable cells (U/u)
            if str(val).strip().upper() == "U":
                u_decisions += 1
                u_postings.add(sid)
                continue

            evaluable_decisions += 1
            human = int(val)
            ext = 1 if skill in extracted else 0

            stats = per_skill[skill]
            if human == 1:
                stats["human_pos"] += 1
            if ext == 1:
                stats["ext_pos"] += 1

            if human == 1 and ext == 1:
                stats["tp"] += 1
            elif human == 0 and ext == 1:
                stats["fp"] += 1
                discrepancies.append({
                    "sample_id": sid,
                    "title": title,
                    "company": company,
                    "skill": skill,
                    "human": human,
                    "extractor": ext,
                    "type": "FP",
                })
            elif human == 1 and ext == 0:
                stats["fn"] += 1
                discrepancies.append({
                    "sample_id": sid,
                    "title": title,
                    "company": company,
                    "skill": skill,
                    "human": human,
                    "extractor": ext,
                    "type": "FN",
                })
            elif human == 0 and ext == 0:
                stats["tn"] += 1

    tot_tp = sum(s["tp"] for s in per_skill.values())
    tot_fp = sum(s["fp"] for s in per_skill.values())
    tot_fn = sum(s["fn"] for s in per_skill.values())
    tot_tn = sum(s["tn"] for s in per_skill.values())
    tot_prec = tot_tp / (tot_tp + tot_fp) if (tot_tp + tot_fp) > 0 else 0.0
    tot_rec = tot_tp / (tot_tp + tot_fn) if (tot_tp + tot_fn) > 0 else 0.0

    return {
        "total_postings": total_postings,
        "evaluable_postings": total_postings - len(u_postings),
        "u_postings_count": len(u_postings),
        "total_decisions": total_decisions,
        "evaluable_decisions": evaluable_decisions,
        "u_decisions": u_decisions,
        "tp": tot_tp,
        "fp": tot_fp,
        "fn": tot_fn,
        "tn": tot_tn,
        "precision": tot_prec,
        "recall": tot_rec,
        "per_skill": per_skill,
        "discrepancies": discrepancies,
    }


def main() -> None:
    # Ensure UTF-8 console output on Windows
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    results = run_validation()

    print("=" * 80)
    print("HW4 SKILL EXTRACTION VALIDATION & DIFF AUDIT")
    print("=" * 80)
    print(f"Sample File:    {SAMPLE_PATH.relative_to(ROOT_DIR)}")
    print(f"Validation File: {WORKBOOK_PATH.relative_to(ROOT_DIR)} (worksheet: Annotation)")
    print(f"Vocabulary File: {VOCAB_PATH.relative_to(ROOT_DIR)}")
    print("-" * 80)

    print("\n[1] SAMPLE & DECISION COUNTS")
    print(f"  Total Sampled Postings:    {results['total_postings']}")
    print(f"  Evaluable Postings:        {results['evaluable_postings']} (Sample 19 excluded due to incomplete text)")
    print(f"  Unevaluable Postings:      {results['u_postings_count']}")
    print(f"  Total Skill Decisions:     {results['total_decisions']} (50 postings × 52 skills)")
    print(f"  Evaluable Skill Decisions: {results['evaluable_decisions']}")
    print(f"  Unevaluable Decisions (U): {results['u_decisions']} (marked U in workbook)")

    print("\n[2] OVERALL CONFUSION MATRIX & ACCURACY")
    human_pos = results["tp"] + results["fn"]
    ext_pos = results["tp"] + results["fp"]
    print(f"  Human Positives:           {human_pos}")
    print(f"  Extractor Positives:       {ext_pos}")
    print(f"  True Positives (TP):       {results['tp']}")
    print(f"  False Positives (FP):      {results['fp']}")
    print(f"  False Negatives (FN):      {results['fn']}")
    print(f"  True Negatives (TN):       {results['tn']}")
    print(f"  Overall Precision:         {results['precision']:.4f} ({results['precision']*100:.2f}%)")
    print(f"  Overall Recall:            {results['recall']:.4f} ({results['recall']*100:.2f}%)")

    print("\n[3] KEY FOCUS SKILLS (R & EXCEL)")
    for sk in ["R", "Excel"]:
        s_data = results["per_skill"][sk]
        s_tp = s_data["tp"]
        s_fp = s_data["fp"]
        s_fn = s_data["fn"]
        s_tn = s_data["tn"]
        s_prec = s_tp / (s_tp + s_fp) if (s_tp + s_fp) > 0 else 0.0
        s_rec = s_tp / (s_tp + s_fn) if (s_tp + s_fn) > 0 else 0.0
        print(f"  * {sk:10s} -> TP: {s_tp:2d} | FP: {s_fp:2d} | FN: {s_fn:2d} | TN: {s_tn:2d} | "
              f"Precision: {s_prec*100:6.2f}% | Recall: {s_rec*100:6.2f}%")

    print("\n[4] PER-SKILL EVALUATION BREAKDOWN (ALL 52 CANONICAL SKILLS)")
    header = f"{'Skill':<25} {'Human+':>7} {'Ext+':>6} {'TP':>5} {'FP':>5} {'FN':>5} {'TN':>5} {'Precision':>10} {'Recall':>10}"
    print(header)
    print("-" * len(header))

    for skill, s_data in results["per_skill"].items():
        s_tp = s_data["tp"]
        s_fp = s_data["fp"]
        s_fn = s_data["fn"]
        s_tn = s_data["tn"]
        s_prec = f"{s_tp / (s_tp + s_fp)*100:6.1f}%" if (s_tp + s_fp) > 0 else "N/A"
        s_rec = f"{s_tp / (s_tp + s_fn)*100:6.1f}%" if (s_tp + s_fn) > 0 else "N/A"
        print(f"{skill:<25} {s_data['human_pos']:>7d} {s_data['ext_pos']:>6d} {s_tp:>5d} {s_fp:>5d} {s_fn:>5d} {s_tn:>5d} {s_prec:>10} {s_rec:>10}")

    print("\n[5] REMAINING DISCREPANCIES (HUMAN != EXTRACTOR)")
    discrepancies = results["discrepancies"]
    if not discrepancies:
        print("  None. Extractor perfectly matches all human annotations.")
    else:
        print(f"  Total remaining discrepancies: {len(discrepancies)}")
        for d in discrepancies:
            print(f"  - Sample {d['sample_id']:2d} [{d['type']}]: Skill='{d['skill']}' | "
                  f"Human={d['human']} vs Extractor={d['extractor']} | "
                  f"Job: '{d['title']}' ({d['company']})")

    print("\n" + "=" * 80)
    print("VERIFICATION RESULT:")
    expected_prec = 0.9884
    expected_rec = 1.0000
    if abs(results["precision"] - expected_prec) < 0.001 and abs(results["recall"] - expected_rec) < 0.0001:
        print("SUCCESS: Exact HW4 validated benchmark reproduced (Precision: 98.84%, Recall: 100.00%).")
    else:
        print("WARNING: Metrics deviate from expected benchmark.")
    print("=" * 80)


if __name__ == "__main__":
    main()
